"""Sample resumes: formats, intended claims, the scanned PDF needing OCR, and
regeneration by scripts/generate_sample_resumes.py."""

from __future__ import annotations

import base64
import importlib.util
import io
import re
import sys
from pathlib import Path
from types import ModuleType

import pytest

from app.seed.loaders import default_samples_dir

RESUMES = default_samples_dir() / "resumes"
# The generator lives in the repository (SAMPLES_DIR may point elsewhere).
SCRIPT = Path(__file__).resolve().parents[5] / "scripts" / "generate_sample_resumes.py"
EXPECTED_FILES = {
    "jane_doe_devops.pdf",
    "jane_doe_devops.docx",
    "jane_doe_devops.txt",
    "jane_doe_devops.png",
    "jane_doe_devops.jpg",
    "jane_doe_devops_scanned.pdf",
    "marcus_okafor_backend.pdf",
    "sofia_martinez_frontend.docx",
    "priya_raman_ml.txt",
}
JANE_CLAIMS = (
    "Designed production Kubernetes platform supporting 10,000 users",
    "Reduced cloud costs by 40%",
    "Led migration of 40 microservices",
    "team of 3",
    "5.8 years",
)
SENSITIVE = re.compile(
    r"date of birth|\bage\b|marital|nationality|citizenship|religion|gender|photo", re.IGNORECASE
)


def _pdf_text(path: Path) -> list[str]:
    import pypdf

    return [page.extract_text() or "" for page in pypdf.PdfReader(str(path)).pages]


def _docx_text(path: Path) -> str:
    import docx

    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)


def _squash(text: str) -> str:
    return re.sub(r"\s+", " ", text)


def test_every_sample_file_exists() -> None:
    assert EXPECTED_FILES <= {p.name for p in RESUMES.iterdir()}


@pytest.mark.parametrize("name", ["jane_doe_devops.pdf", "jane_doe_devops.docx", "jane_doe_devops.txt"])
def test_jane_text_formats_contain_the_claims(name: str) -> None:
    path = RESUMES / name
    if path.suffix == ".pdf":
        text = "\n".join(_pdf_text(path))
    elif path.suffix == ".docx":
        text = _docx_text(path)
    else:
        text = path.read_text(encoding="utf-8")
    flat = _squash(text)
    for claim in JANE_CLAIMS:
        assert claim in flat, f"{name} misses {claim!r}"
    # The deliberate inconsistency: Bluepeak starts before Harborline ends.
    assert "Jan 2023 – Present" in flat and "Aug 2021 – Mar 2023" in flat
    assert not SENSITIVE.search(flat)


def test_scanned_pdf_is_image_only() -> None:
    import pypdf

    reader = pypdf.PdfReader(str(RESUMES / "jane_doe_devops_scanned.pdf"))
    assert len(reader.pages) >= 2
    for page in reader.pages:
        assert (page.extract_text() or "").strip() == ""
        assert len(page.images) == 1
        image = page.images[0].image
        assert image is not None and image.width >= 1200


def test_image_resume_is_ocr_sized() -> None:
    from PIL import Image

    for name in ("jane_doe_devops.png", "jane_doe_devops.jpg"):
        with Image.open(RESUMES / name) as image:
            assert image.width >= 1200 and image.height >= 1600


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        (
            "marcus_okafor_backend.pdf",
            ("Backend Engineer (Python)", "1.2 million requests per day", "900 ms to 180 ms"),
        ),
        (
            "sofia_martinez_frontend.docx",
            ("Junior Frontend Engineer", "from 72 to 96", "React Testing Library"),
        ),
        ("priya_raman_ml.txt", ("Machine Learning Engineer (LLM)", "61% to 78%", "1.4 s to 380 ms")),
    ],
)
def test_other_candidates(name: str, expected: tuple[str, ...]) -> None:
    path = RESUMES / name
    text = (
        "\n".join(_pdf_text(path))
        if path.suffix == ".pdf"
        else _docx_text(path)
        if path.suffix == ".docx"
        else path.read_text(encoding="utf-8")
    )
    flat = _squash(text)
    for needle in expected:
        assert needle in flat, f"{name} misses {needle!r}"
    assert not SENSITIVE.search(flat)
    assert "@example.com" in flat


def test_readme_documents_every_file() -> None:
    readme = (RESUMES / "README.md").read_text(encoding="utf-8")
    for name in EXPECTED_FILES:
        assert f"`{name}`" in readme


def _load_script() -> ModuleType:
    spec = importlib.util.spec_from_file_location("generate_sample_resumes", SCRIPT)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # dataclasses resolve annotations through sys.modules
    try:
        spec.loader.exec_module(module)
    except BaseException:
        sys.modules.pop(spec.name, None)
        raise
    return module


def test_generator_reproduces_the_samples(tmp_path: Path) -> None:
    script = _load_script()
    written = {p.name for p in script.generate(tmp_path)}
    assert written == EXPECTED_FILES
    for name in ("jane_doe_devops.txt", "priya_raman_ml.txt"):
        assert (tmp_path / name).read_bytes() == (RESUMES / name).read_bytes()
    for name in ("jane_doe_devops.pdf", "marcus_okafor_backend.pdf"):
        assert _pdf_text(tmp_path / name) == _pdf_text(RESUMES / name)
    assert (tmp_path / "jane_doe_devops.pdf").read_bytes() == (RESUMES / "jane_doe_devops.pdf").read_bytes()


@pytest.mark.live
def test_ocr_endpoint_reads_the_scanned_pdf() -> None:
    """One live call to the OCR model with page 1 of the scanned PDF."""
    import httpx
    import pypdfium2 as pdfium

    from app.core.config import get_settings

    settings = get_settings()
    page = pdfium.PdfDocument(str(RESUMES / "jane_doe_devops_scanned.pdf"))[0]
    image = page.render(scale=150 / 72).to_pil().convert("L")
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    content = [
        {
            "type": "image_url",
            "image_url": {"url": "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()},
        },
        {"type": "text", "text": settings.ocr_prompt},
    ]
    response = httpx.post(
        f"{settings.ocr_base_url}/chat/completions",
        json={
            "model": settings.ocr_model,
            "messages": [{"role": "user", "content": content}],
            "max_tokens": 2048,
            "temperature": 0,
        },
        timeout=settings.ocr_timeout_seconds,
    )
    response.raise_for_status()
    text = response.json()["choices"][0]["message"]["content"]
    for needle in ("Jane Doe", "10,000", "40 microservices", "40%", "Mar 2023"):
        assert needle in text
