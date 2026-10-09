"""Extraction for every sample resume format (OCR mocked)."""

from __future__ import annotations

import io

import pytest
from PIL import Image

from app.ai.base import AIUnavailable, OcrResult
from app.domain.enums import ExtractionMethod
from app.services.resume.extraction import (
    NoReadableText,
    ResumeExtractor,
    encode_for_ocr,
    kind_from_content_type,
    prepare_image_pages,
    render_pdf_pages,
)
from app.services.resume.normalize import detect_sections, normalize_text
from tests.unit.resume import corpus

JANE_TXT = corpus.sample("jane_doe_devops.txt").decode()
PAGE_1, PAGE_2 = (
    JANE_TXT[: JANE_TXT.index("DevOps Engineer — Harborline")],
    JANE_TXT[JANE_TXT.index("DevOps Engineer — Harborline") :],
)


def extractor(mock_ai, max_pages: int = 12) -> ResumeExtractor:
    return ResumeExtractor(mock_ai, ocr_max_pages=max_pages, organization_id="org")


def assert_jane(text: str) -> None:
    normalized = normalize_text(text)
    for needle in ("Bluepeak Analytics", "Harborline Logistics", "10,000 users", "40%", "Terraform"):
        assert needle in normalized
    assert [s.key for s in detect_sections(normalized)][:3] == ["other", "summary", "experience"]


async def test_text_pdf_without_ocr(mock_ai) -> None:
    result = await extractor(mock_ai).extract(corpus.sample("jane_doe_devops.pdf"), "pdf")
    assert result.method is ExtractionMethod.TEXT
    assert result.page_count == 1 and result.ocr_pages == 0
    assert mock_ai.ocr.calls == []
    assert_jane(result.text)


async def test_docx(mock_ai) -> None:
    result = await extractor(mock_ai).extract(corpus.sample("jane_doe_devops.docx"), "docx")
    assert result.method is ExtractionMethod.TEXT and result.page_count is None
    assert "- Designed production Kubernetes platform" in normalize_text(result.text)
    assert_jane(result.text)


async def test_docx_tables_headers_and_text_boxes(mock_ai) -> None:
    import docx

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Alex Rivera | alex@example.com"
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("Platform Engineer at Northwind", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Years"
    table.cell(1, 0).text = "Kubernetes"
    table.cell(1, 1).text = "5"
    document.sections[0].footer.paragraphs[0].text = "References available on request"
    buffer = io.BytesIO()
    document.save(buffer)
    result = await extractor(mock_ai).extract(buffer.getvalue(), "docx")
    lines = normalize_text(result.text).split("\n")
    assert lines[0] == "Alex Rivera | alex@example.com"
    assert "- Platform Engineer at Northwind" in lines
    assert "Skill | Years" in lines and "Kubernetes | 5" in lines
    assert lines[-1] == "References available on request"


async def test_txt(mock_ai) -> None:
    result = await extractor(mock_ai).extract(corpus.sample("jane_doe_devops.txt"), "txt")
    assert result.method is ExtractionMethod.TEXT
    assert_jane(result.text)
    latin_text = "Jos\u00e9 \u00c1lvarez\nIngeniero de plataformas con 6 a\u00f1os de experiencia en Kubernetes y Terraform."
    latin = await extractor(mock_ai).extract(latin_text.encode("latin-1"), "txt")
    assert latin.text == latin_text and latin.warnings == ["encoding_latin-1"]


@pytest.mark.parametrize(("name", "kind"), [("jane_doe_devops.png", "png"), ("jane_doe_devops.jpg", "jpeg")])
async def test_image_resume_uses_ocr(mock_ai, name: str, kind: str) -> None:
    mock_ai.ocr.script(JANE_TXT)
    progress: list[tuple[int, int]] = []

    async def on_page(done: int, total: int) -> None:
        progress.append((done, total))

    result = await extractor(mock_ai).extract(corpus.sample(name), kind, on_ocr_page=on_page)
    assert result.method is ExtractionMethod.OCR
    assert result.ocr_pages == 1 and result.page_count == 1
    assert progress == [(1, 1)]
    call = mock_ai.ocr.calls[0]
    assert call["mime_type"] in {"image/png", "image/jpeg"}
    assert call["bytes"] < 2_000_000
    assert_jane(result.text)


async def test_scanned_pdf_uses_ocr_per_page(mock_ai) -> None:
    mock_ai.ocr.script(PAGE_1, PAGE_2)
    started: list[int] = []
    pages: list[tuple[int, int]] = []

    async def on_start(_: int, total: int) -> None:
        started.append(total)

    async def on_page(done: int, total: int) -> None:
        pages.append((done, total))

    result = await extractor(mock_ai).extract(
        corpus.sample("jane_doe_devops_scanned.pdf"), "pdf", on_ocr_start=on_start, on_ocr_page=on_page
    )
    assert result.method is ExtractionMethod.OCR
    assert result.page_count == 2 and result.ocr_pages == 2
    assert started == [2] and pages[-1] == (2, 2)
    assert [p.method for p in result.pages] == ["ocr", "ocr"]
    assert len(mock_ai.ocr.calls) == 2
    assert_jane(result.text)


async def test_scanned_pdf_respects_page_limit(mock_ai) -> None:
    mock_ai.ocr.script(JANE_TXT)
    result = await extractor(mock_ai, max_pages=1).extract(
        corpus.sample("jane_doe_devops_scanned.pdf"), "pdf"
    )
    assert result.ocr_pages == 1
    assert "ocr_page_limit:1" in result.warnings
    assert len(mock_ai.ocr.calls) == 1


async def test_mixed_pdf(mock_ai) -> None:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(corpus.sample("jane_doe_devops.pdf"))))
    writer.append(PdfReader(io.BytesIO(corpus.sample("jane_doe_devops_scanned.pdf"))), pages=[0])
    buffer = io.BytesIO()
    writer.write(buffer)
    mock_ai.ocr.script("Scanned certificate page: Certified Kubernetes Administrator (CKA), issued 2023.")
    result = await extractor(mock_ai).extract(buffer.getvalue(), "pdf")
    assert result.method is ExtractionMethod.MIXED
    assert [p.method for p in result.pages] == ["text", "ocr"]


async def test_ocr_page_error_is_a_warning(mock_ai) -> None:
    mock_ai.ocr.script(OcrResult(text="", status="error"), PAGE_2)
    result = await extractor(mock_ai).extract(corpus.sample("jane_doe_devops_scanned.pdf"), "pdf")
    assert any(w.startswith("ocr_failed_page:") for w in result.warnings)
    assert "Harborline" in result.text


async def test_ocr_outage_propagates_for_retry(mock_ai) -> None:
    mock_ai.ocr.script(AIUnavailable("ocr down", endpoint="ocr"), AIUnavailable("ocr down", endpoint="ocr"))
    with pytest.raises(AIUnavailable):
        await extractor(mock_ai).extract(corpus.sample("jane_doe_devops_scanned.pdf"), "pdf")


async def test_blank_image_has_no_readable_text(mock_ai) -> None:
    mock_ai.ocr.script("")
    with pytest.raises(NoReadableText):
        await extractor(mock_ai).extract(corpus.png_bytes(), "png")


def test_rendering_and_encoding() -> None:
    images = render_pdf_pages(corpus.sample("jane_doe_devops_scanned.pdf"), [0, 1], dpi=200)
    assert set(images) == {0, 1}
    for payload, mime in images.values():
        assert mime in {"image/png", "image/jpeg"}
        with Image.open(io.BytesIO(payload)) as image:
            assert image.mode == "L"
            assert max(image.size) <= 3500 and min(image.size) >= 1500  # ~200 DPI letter
    clean, mime = encode_for_ocr(Image.new("RGB", (5000, 300), "white"))
    assert mime == "image/png"
    with Image.open(io.BytesIO(clean)) as image:
        assert image.size[0] == 3500


def test_exif_orientation_is_fixed() -> None:
    source = Image.new("RGB", (400, 1000), "white")
    exif = source.getexif()
    exif[0x0112] = 6  # rotate 90 degrees when displayed
    buffer = io.BytesIO()
    source.save(buffer, format="JPEG", exif=exif)
    [(payload, _)], frames = prepare_image_pages(buffer.getvalue(), 5)
    assert frames == 1
    with Image.open(io.BytesIO(payload)) as image:
        assert image.size == (1000, 400)


def test_kind_from_content_type() -> None:
    assert kind_from_content_type("application/pdf") == "pdf"
    assert kind_from_content_type("image/jpeg; q=1") == "jpeg"
    with pytest.raises(ValueError):
        kind_from_content_type("application/x-msdownload")
