"""Text extraction from validated resume files, with OCR where needed.

* **PDF** - ``pdfplumber`` text per page (``pypdf`` per-page fallback). A page
  with fewer than ``MIN_PAGE_CHARS`` meaningful characters is rendered with
  ``pypdfium2`` at ~200 DPI and sent to the OCR model (at most
  ``OCR_MAX_PAGES`` pages, page 1 first, several at a time; the readable
  prefix is reported while later pages are scanned). The OCR text replaces the text
  layer only when it is richer.
* **DOCX** - ``python-docx``: body paragraphs and tables in document order
  (layout tables cell by cell), headers/footers, text boxes; list paragraphs
  become ``- `` bullets.
* **TXT** - decoded like the validator does.
* **Images** - EXIF orientation fixed, downscaled to a sane size, every frame
  (TIFF) up to the page limit OCR'd.

``ExtractionMethod`` is ``text`` (only text layers), ``ocr`` (only OCR) or
``mixed``. OCR outages (``AIUnavailable``/``AITimeout``) propagate so the job
retries; an OCR rejection of one page (``status="error"``) only adds a
warning.
"""

from __future__ import annotations

import asyncio
import io
import threading
import time
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field
from typing import Any, Literal

from PIL import Image, ImageOps

from app.ai.base import AIServices, CallContext
from app.core.logging import get_logger
from app.domain.enums import ExtractionMethod, ModelTask
from app.services.resume.normalize import meaningful_chars
from app.services.resume.validation import IMAGE_KINDS, decode_text

log = get_logger(__name__)

MIN_PAGE_CHARS = 80
OCR_DPI = 200
# Pages OCR'd at once (the OCR endpoint client caps the process-wide concurrency;
# measured 2026-10-08 on the OCR server, one 1,055-token page: 1 request 15.8 s,
# 2 at once 19.7 s each, 3 23.7 s, 4 27.9 s - aggregate 67 -> 151 tok/s).
OCR_CONCURRENCY = 4
MAX_RENDER_SIDE = 3500  # px; guards against absurd page sizes
MAX_EXTRACTED_CHARS = 200_000

# pdfium is not thread-safe: serialise every call across the process.
_PDFIUM_LOCK = threading.Lock()

ProgressCallback = Callable[[int, int], Awaitable[None]]
# (text of pages 1..n in order, n, total pages): called whenever the readable
# prefix of the document grows while pages are still being scanned.
PartialTextCallback = Callable[[str, int, int], Awaitable[None]]


class NoReadableText(Exception):
    """The file holds no readable text, even after OCR."""


@dataclass(slots=True)
class PageResult:
    number: int  # 1-based
    method: Literal["text", "ocr", "none"]
    chars: int
    text: str = field(repr=False)
    ocr_status: str | None = None
    ocr_ms: float | None = None


@dataclass(slots=True)
class ExtractionResult:
    text: str
    method: ExtractionMethod
    page_count: int | None
    pages: list[PageResult] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    ocr_pages: int = 0
    ocr_ms: float = 0.0
    duration_ms: float = 0.0

    def summary(self) -> dict[str, Any]:
        return {
            "method": self.method.value,
            "page_count": self.page_count,
            "ocr_pages": self.ocr_pages,
            "ocr_ms": round(self.ocr_ms, 1),
            "duration_ms": round(self.duration_ms, 1),
            "characters": len(self.text),
            "pages": [
                {"page": p.number, "method": p.method, "chars": p.chars, "ocr_status": p.ocr_status}
                for p in self.pages[:50]
            ],
            "warnings": list(self.warnings),
        }


def kind_from_content_type(content_type: str) -> str:
    mapping = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "text/plain": "txt",
        "image/png": "png",
        "image/jpeg": "jpeg",
        "image/webp": "webp",
        "image/tiff": "tiff",
    }
    try:
        return mapping[content_type.split(";")[0].strip().lower()]
    except KeyError as exc:
        raise ValueError(f"unsupported stored content type {content_type!r}") from exc


class ResumeExtractor:
    def __init__(
        self,
        ai: AIServices,
        *,
        ocr_max_pages: int,
        organization_id: str | None = None,
        ocr_concurrency: int = OCR_CONCURRENCY,
        dpi: int = OCR_DPI,
    ) -> None:
        self.ai = ai
        self.ocr_max_pages = max(0, ocr_max_pages)
        self.organization_id = organization_id
        self.semaphore = asyncio.Semaphore(max(1, ocr_concurrency))
        self.dpi = dpi
        self._on_partial_text: PartialTextCallback | None = None

    async def extract(
        self,
        data: bytes,
        kind: str,
        *,
        on_ocr_start: ProgressCallback | None = None,
        on_ocr_page: ProgressCallback | None = None,
        on_partial_text: PartialTextCallback | None = None,
    ) -> ExtractionResult:
        started = time.perf_counter()
        self._on_partial_text = on_partial_text
        if kind == "pdf":
            result = await self._pdf(data, on_ocr_start, on_ocr_page)
        elif kind == "docx":
            text = await asyncio.to_thread(docx_text, data)
            result = ExtractionResult(text=text, method=ExtractionMethod.TEXT, page_count=None)
        elif kind == "txt":
            text, encoding = decode_text(data)
            result = ExtractionResult(text=text, method=ExtractionMethod.TEXT, page_count=None)
            if encoding not in ("utf-8", "utf-8-sig"):
                result.warnings.append(f"encoding_{encoding}")
        elif kind in IMAGE_KINDS:
            result = await self._image(data, on_ocr_start, on_ocr_page)
        else:
            raise ValueError(f"unsupported kind {kind!r}")
        if len(result.text) > MAX_EXTRACTED_CHARS:
            result.text = result.text[:MAX_EXTRACTED_CHARS]
            result.warnings.append("text_truncated")
        result.duration_ms = (time.perf_counter() - started) * 1000
        if meaningful_chars(result.text) < 40:
            raise NoReadableText(kind)
        return result

    # ------------------------------------------------------------------ OCR
    async def _ocr(self, image: OcrImage, page_number: int) -> tuple[str, str, float]:
        payload, mime_type = image
        async with self.semaphore:
            started = time.perf_counter()
            result = await self.ai.ocr.ocr_image(
                payload,
                mime_type=mime_type,
                context=CallContext(task=ModelTask.OCR, organization_id=self.organization_id),
            )
            elapsed = (time.perf_counter() - started) * 1000
        log.info(
            "resume.ocr_page",
            page=page_number,
            status=result.status,
            chars=len(result.text),
            ms=round(elapsed),
        )
        text = result.text if result.status in ("ok", "degenerate") else ""
        return text, result.status, elapsed

    async def _ocr_pages(
        self,
        images: dict[int, OcrImage],
        on_ocr_page: ProgressCallback | None,
        *,
        known: dict[int, str] | None = None,
        page_count: int | None = None,
    ) -> dict[int, tuple[str, str, float]]:
        """OCR the pages concurrently (page 1 first). ``known`` holds the text of
        pages that need no OCR; whenever the readable prefix (pages 1..n) grows
        before the last page is done, ``on_partial_text`` gets that prefix so the
        candidate sees a first read of page 1 while page 2 is still scanned."""
        total = len(images)
        pages_total = page_count or total
        done = 0
        results: dict[int, tuple[str, str, float]] = {}
        texts: dict[int, str] = dict(known or {})
        prefix_pages = 0
        lock = asyncio.Lock()
        while prefix_pages + 1 in texts:  # leading pages with a text layer: show them right away
            prefix_pages += 1
        if prefix_pages and self._on_partial_text is not None and prefix_pages < pages_total:
            joined = "\n\n".join(texts[i].strip() for i in range(1, prefix_pages + 1) if texts[i].strip())
            try:
                await self._on_partial_text(joined, prefix_pages, pages_total)
            except Exception as exc:  # noqa: BLE001
                log.warning("resume.partial_text_failed", error=type(exc).__name__)

        async def run(number: int, image: OcrImage) -> None:
            nonlocal done, prefix_pages
            results[number] = await self._ocr(image, number)
            async with lock:
                done += 1
                text, status, _ = results[number]
                previous = texts.get(number, "")
                texts[number] = text if meaningful_chars(text) > meaningful_chars(previous) else previous
                if on_ocr_page is not None:
                    await on_ocr_page(done, total)
                prefix = 0
                while prefix + 1 in texts:
                    prefix += 1
                if self._on_partial_text is not None and prefix > prefix_pages and prefix < pages_total:
                    prefix_pages = prefix
                    joined = "\n\n".join(texts[i].strip() for i in range(1, prefix + 1) if texts[i].strip())
                    try:
                        await self._on_partial_text(joined, prefix, pages_total)
                    except Exception as exc:  # noqa: BLE001 - a progress update must never fail the OCR
                        log.warning("resume.partial_text_failed", error=type(exc).__name__)

        await asyncio.gather(*(run(number, image) for number, image in sorted(images.items())))
        return results

    # ------------------------------------------------------------------ PDF
    async def _pdf(
        self, data: bytes, on_ocr_start: ProgressCallback | None, on_ocr_page: ProgressCallback | None
    ) -> ExtractionResult:
        texts = await asyncio.to_thread(pdf_text_pages, data)
        pages = [
            PageResult(
                number=i + 1,
                method="text" if meaningful_chars(t) else "none",
                chars=meaningful_chars(t),
                text=t,
            )
            for i, t in enumerate(texts)
        ]
        warnings: list[str] = []
        needs_ocr = [p.number for p in pages if p.chars < MIN_PAGE_CHARS]
        selected = needs_ocr[: self.ocr_max_pages]
        if len(needs_ocr) > len(selected):
            warnings.append(f"ocr_page_limit:{len(needs_ocr) - len(selected)}")
        ocr_ms = 0.0
        if selected:
            if on_ocr_start is not None:
                await on_ocr_start(0, len(selected))
            images = await asyncio.to_thread(render_pdf_pages, data, [n - 1 for n in selected], self.dpi)
            known = {p.number: p.text for p in pages if p.number not in set(selected)}
            results = await self._ocr_pages(
                {n + 1: img for n, img in images.items()}, on_ocr_page, known=known, page_count=len(pages)
            )
            for number, (text, status, elapsed) in results.items():
                page = pages[number - 1]
                page.ocr_status, page.ocr_ms = status, round(elapsed, 1)
                ocr_ms += elapsed
                if status == "error":
                    warnings.append(f"ocr_failed_page:{number}")
                elif status == "degenerate":
                    warnings.append(f"ocr_degenerate_page:{number}")
                chars = meaningful_chars(text)
                if chars > page.chars:
                    page.text, page.chars, page.method = text, chars, "ocr"
        used = {p.method for p in pages if p.chars > 0}
        if used == {"ocr"}:
            method = ExtractionMethod.OCR
        elif "ocr" in used:
            method = ExtractionMethod.MIXED
        else:
            method = ExtractionMethod.TEXT
        text = "\n\n".join(p.text.strip() for p in pages if p.text.strip())
        return ExtractionResult(
            text=text,
            method=method,
            page_count=len(pages),
            pages=pages,
            warnings=warnings,
            ocr_pages=len(selected),
            ocr_ms=ocr_ms,
        )

    # --------------------------------------------------------------- images
    async def _image(
        self, data: bytes, on_ocr_start: ProgressCallback | None, on_ocr_page: ProgressCallback | None
    ) -> ExtractionResult:
        limit = max(1, self.ocr_max_pages)
        images, frames = await asyncio.to_thread(prepare_image_pages, data, limit)
        warnings = [f"ocr_page_limit:{frames - len(images)}"] if frames > len(images) else []
        if on_ocr_start is not None:
            await on_ocr_start(0, len(images))
        results = await self._ocr_pages(
            {i + 1: img for i, img in enumerate(images)}, on_ocr_page, page_count=len(images)
        )
        pages: list[PageResult] = []
        ocr_ms = 0.0
        for number in sorted(results):
            text, status, elapsed = results[number]
            ocr_ms += elapsed
            if status == "error":
                warnings.append(f"ocr_failed_page:{number}")
            elif status == "degenerate":
                warnings.append(f"ocr_degenerate_page:{number}")
            chars = meaningful_chars(text)
            pages.append(
                PageResult(
                    number,
                    "ocr" if chars else "none",
                    chars,
                    text,
                    ocr_status=status,
                    ocr_ms=round(elapsed, 1),
                )
            )
        return ExtractionResult(
            text="\n\n".join(p.text.strip() for p in pages if p.text.strip()),
            method=ExtractionMethod.OCR,
            page_count=frames,
            pages=pages,
            warnings=warnings,
            ocr_pages=len(images),
            ocr_ms=ocr_ms,
        )


# =============================================================================
# Blocking helpers (run in threads)
# =============================================================================


def pdf_text_pages(data: bytes) -> list[str]:
    """Text of every page: pdfplumber first, pypdf for pages it cannot read."""
    texts: list[str | None] = []
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data), password="") as pdf:
            for page in pdf.pages:
                try:
                    texts.append(page.extract_text(x_tolerance=1.5, y_tolerance=3) or "")
                except Exception as exc:  # noqa: BLE001 - one broken page must not lose the others
                    log.info("resume.pdfplumber_page_failed", error=type(exc).__name__)
                    texts.append(None)
                finally:
                    page.close()
    except Exception as exc:  # noqa: BLE001
        log.info("resume.pdfplumber_failed", error=type(exc).__name__)
        texts = []
    if texts and all(t is not None for t in texts):
        return [t or "" for t in texts]
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=False)
    if reader.is_encrypted:
        reader.decrypt("")
    out: list[str] = []
    for index, page in enumerate(reader.pages):
        if index < len(texts) and texts[index] is not None:
            out.append(texts[index] or "")
            continue
        try:
            out.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            log.info("resume.pypdf_page_failed", page=index + 1, error=type(exc).__name__)
            out.append("")
    return out


OcrImage = tuple[bytes, str]  # (encoded image, mime type)
PNG_MAX_BYTES = 1_500_000  # noisy scans compress badly as PNG: JPEG q90 is ~2.3x smaller and 30x faster


def encode_for_ocr(image: Image.Image) -> OcrImage:
    """Greyscale (text recognition does not need colour), capped size, PNG for
    clean renderings and JPEG for noisy scans/photos."""
    if max(image.size) > MAX_RENDER_SIDE:
        image = image.copy()
        image.thumbnail((MAX_RENDER_SIDE, MAX_RENDER_SIDE), Image.Resampling.LANCZOS)
    if image.mode != "L":
        image = image.convert("L")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    if buffer.tell() <= PNG_MAX_BYTES:
        return buffer.getvalue(), "image/png"
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=90)
    return buffer.getvalue(), "image/jpeg"


def render_pdf_pages(data: bytes, indexes: list[int], dpi: int = OCR_DPI) -> dict[int, OcrImage]:
    """Greyscale renderings (0-based page index -> encoded image) at ``dpi``,
    capped at ``MAX_RENDER_SIDE`` pixels per side."""
    import pypdfium2 as pdfium

    out: dict[int, OcrImage] = {}
    with _PDFIUM_LOCK:
        document = pdfium.PdfDocument(data, password="")
        try:
            for index in indexes:
                page = document[index]
                try:
                    width, height = page.get_size()
                    scale = min(dpi / 72.0, MAX_RENDER_SIDE / max(width, height, 1.0))
                    bitmap = page.render(scale=scale, grayscale=True)
                    try:
                        out[index] = encode_for_ocr(bitmap.to_pil())
                    finally:
                        bitmap.close()
                finally:
                    page.close()
        finally:
            document.close()
    return out


def prepare_image_pages(data: bytes, max_pages: int) -> tuple[list[OcrImage], int]:
    """EXIF-corrected, OCR-encoded frames (up to ``max_pages``) and the frame count."""
    images: list[OcrImage] = []
    with Image.open(io.BytesIO(data)) as source:
        frames = int(getattr(source, "n_frames", 1) or 1)
        for index in range(min(frames, max_pages)):
            if frames > 1:
                source.seek(index)
            frame = ImageOps.exif_transpose(source.copy())
            images.append(encode_for_ocr(frame))
    return images, frames


def _is_list_paragraph(paragraph: Any) -> bool:
    p_pr = paragraph._p.pPr  # noqa: SLF001 - python-docx exposes numbering only via the XML
    if p_pr is not None and p_pr.numPr is not None:
        return True
    style = getattr(paragraph.style, "name", "") or ""
    return style.lower().startswith("list")


def _blocks_text(blocks: Any, depth: int = 0) -> list[str]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []
    for block in blocks:
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                lines.append(f"- {text}" if _is_list_paragraph(block) and not text.startswith("-") else text)
            else:
                lines.append("")
        elif isinstance(block, Table) and depth < 4:
            lines.extend(_table_text(block, depth))
    return lines


def _table_text(table: Any, depth: int) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        cells: list[Any] = []
        seen: set[int] = set()
        for cell in row.cells:
            key = id(cell._tc)  # noqa: SLF001 - merged cells repeat the same element
            if key in seen:
                continue
            seen.add(key)
            cells.append(cell)
        texts = [cell.text.strip() for cell in cells]
        layout = any("\n" in t for t in texts) or any(cell.tables for cell in cells)
        if layout:
            for cell in cells:
                block = _blocks_text(cell.iter_inner_content(), depth + 1)
                if any(line.strip() for line in block):
                    lines.extend(block)
                    lines.append("")
        else:
            row_text = " | ".join(t for t in texts if t)
            if row_text:
                lines.append(row_text)
    lines.append("")
    return lines


def _textbox_text(element: Any) -> list[str]:
    from docx.oxml.ns import qn

    seen: set[str] = set()
    blocks: list[str] = []
    for box in element.iter(qn("w:txbxContent")):
        paragraphs = []
        for paragraph in box.iter(qn("w:p")):
            text = "".join(node.text or "" for node in paragraph.iter(qn("w:t"))).strip()
            if text:
                paragraphs.append(text)
        block = "\n".join(paragraphs)
        if block and block not in seen:
            seen.add(block)
            blocks.append(block)
    return blocks


def docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    seen_parts: set[str] = set()
    header_lines: list[str] = []
    footer_lines: list[str] = []

    def collect(part: Any, target: list[str]) -> None:
        if part is None or part.is_linked_to_previous:
            return
        text = "\n".join(_blocks_text(part.iter_inner_content())).strip()
        if text and text not in seen_parts:
            seen_parts.add(text)
            target.append(text)

    for section in document.sections:
        collect(section.header, header_lines)
        collect(section.footer, footer_lines)
        if section.different_first_page_header_footer:
            collect(section.first_page_header, header_lines)
            collect(section.first_page_footer, footer_lines)
    body = _blocks_text(document.iter_inner_content())
    boxes = _textbox_text(document.element.body)
    parts = [*header_lines, "\n".join(body), *boxes, *footer_lines]
    return "\n\n".join(part.strip() for part in parts if part and part.strip())
