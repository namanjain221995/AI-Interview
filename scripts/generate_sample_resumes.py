#!/usr/bin/env python3
"""Generate the realistic sample resumes in ``samples/resumes/``.

Run it with the API virtualenv, which provides reportlab, python-docx, Pillow
and numpy::

    cd apps/api && uv run python ../../scripts/generate_sample_resumes.py [--out DIR]

Files produced (see samples/resumes/README.md for the intended claims and gaps):

* ``jane_doe_devops.pdf``          text-layer PDF (reportlab)
* ``jane_doe_devops.docx``         Word document (python-docx)
* ``jane_doe_devops.txt``          plain text
* ``jane_doe_devops.png`` / ``.jpg``  image resume (Pillow)
* ``jane_doe_devops_scanned.pdf``  image-only pages, slightly rotated and noisy (needs OCR)
* ``marcus_okafor_backend.pdf``    mid-level backend engineer (text-layer PDF)
* ``sofia_martinez_frontend.docx`` junior frontend engineer
* ``priya_raman_ml.txt``           ML engineer (LLM)

Every person, company, school, e-mail address and phone number is fictional
(``example.com`` addresses, 555 phone numbers). Resumes deliberately contain no
photo, date of birth, age, marital status, nationality or other protected
characteristics. Output is deterministic: fixed seeds, fixed document dates and
reportlab's invariant mode.
"""

from __future__ import annotations

import argparse
import io
import sys
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = REPO_ROOT / "samples" / "resumes"
DOC_DATE = datetime(2026, 10, 1, 9, 0, 0, tzinfo=UTC)
DPI = 150
PAGE_W, PAGE_H = int(8.5 * DPI), int(11 * DPI)
MARGIN = 95

try:
    import numpy as np
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer
except ImportError as exc:  # pragma: no cover - guidance for the operator
    sys.exit(
        f"missing dependency ({exc.name}); run with: cd apps/api && uv run python ../../scripts/{Path(__file__).name}"
    )


# =============================================================================
# Resume content
# =============================================================================


@dataclass(frozen=True)
class Role:
    title: str
    company: str
    location: str
    start: str
    end: str
    bullets: tuple[str, ...]


@dataclass(frozen=True)
class Resume:
    slug: str
    name: str
    headline: str
    location: str
    email: str
    phone: str
    website: str | None
    summary: str
    roles: tuple[Role, ...]
    skills: tuple[tuple[str, str], ...]
    projects: tuple[tuple[str, str], ...] = ()
    certifications: tuple[str, ...] = ()
    education: tuple[str, ...] = ()
    extra_sections: tuple[tuple[str, tuple[str, ...]], ...] = field(default_factory=tuple)

    @property
    def contact_line(self) -> str:
        parts = [self.location, self.email, self.phone]
        if self.website:
            parts.append(self.website)
        return " · ".join(parts)


JANE = Resume(
    slug="jane_doe_devops",
    name="Jane Doe",
    headline="Senior DevOps Engineer",
    location="Austin, TX",
    email="jane.doe@example.com",
    phone="+1 (512) 555-0147",
    website="janedoe.example.com",
    summary=(
        "Senior DevOps engineer with 5.8 years of experience building and operating cloud-native platforms on AWS "
        "and Kubernetes. Designed a production Kubernetes platform supporting 10,000 users, led the migration of 40 "
        "microservices to Amazon EKS and reduced cloud costs by 40%. Lead a team of 3 engineers and care about "
        "reliability, automation and developer experience."
    ),
    roles=(
        Role(
            "Senior DevOps Engineer",
            "Bluepeak Analytics",
            "Austin, TX",
            "Jan 2023",
            "Present",
            (
                "Designed production Kubernetes platform supporting 10,000 users on Amazon EKS (3 clusters, about 120 nodes).",
                "Led migration of 40 microservices from EC2 and ECS to EKS over nine months with no customer-facing downtime.",
                "Reduced cloud costs by 40% through rightsizing, Karpenter-managed spot capacity and storage lifecycle policies.",
                "Lead a team of 3 DevOps engineers: sprint planning, code reviews, on-call rotation and onboarding.",
                "Introduced GitOps with Argo CD and progressive delivery with Argo Rollouts; deployment lead time fell from 2 days to 45 minutes.",
                "Built observability with Prometheus, Grafana, Loki and OpenTelemetry; defined SLOs for 12 customer-facing services.",
            ),
        ),
        Role(
            "DevOps Engineer",
            "Harborline Logistics",
            "Remote",
            "Aug 2021",
            "Mar 2023",
            (
                "Managed AWS infrastructure across 4 accounts with Terraform (60+ modules) and Terragrunt.",
                "Built GitHub Actions CI/CD pipelines for 25 services and cut average build time by 55%.",
                "Served in the on-call rotation; led 14 incident postmortems and introduced a blameless postmortem template.",
                "Containerised legacy Java services and standardised deployments with Helm charts.",
            ),
        ),
        Role(
            "Junior Systems Engineer",
            "Cobalt Retail Group",
            "Dallas, TX",
            "Jan 2021",
            "Jul 2021",
            (
                "Automated provisioning of a 200-VM Linux fleet with Ansible.",
                "Wrote Python tooling for log rotation checks and nightly backup verification.",
            ),
        ),
    ),
    skills=(
        (
            "Cloud & IaC",
            "AWS (EKS, EC2, VPC, IAM, RDS, CloudFront, Organizations), Terraform, Terragrunt, Ansible",
        ),
        ("Containers", "Kubernetes, Helm, Karpenter, Argo CD, Argo Rollouts, Docker"),
        ("CI/CD", "GitHub Actions, GitLab CI, Jenkins"),
        ("Observability", "Prometheus, Grafana, Loki, OpenTelemetry, PagerDuty"),
        ("Languages", "Python, Go, Bash"),
        ("Data stores", "PostgreSQL, Redis"),
    ),
    projects=(
        (
            "Internal developer platform",
            "Self-service namespaces and golden-path service templates in Backstage, adopted by 8 product teams.",
        ),
    ),
    certifications=(
        "AWS Certified Solutions Architect – Associate (2022)",
        "Certified Kubernetes Administrator (CKA) (2023)",
    ),
    education=("B.S. Computer Science, Lakeview State University",),
)

MARCUS = Resume(
    slug="marcus_okafor_backend",
    name="Marcus Okafor",
    headline="Backend Engineer (Python)",
    location="Chicago, IL",
    email="marcus.okafor@example.com",
    phone="+1 (312) 555-0193",
    website=None,
    summary=(
        "Backend engineer with 3.5 years of experience building Python APIs on PostgreSQL. Focused on correctness, "
        "testing and fast queries; comfortable owning services from design to production."
    ),
    roles=(
        Role(
            "Backend Engineer",
            "Fieldnote Health",
            "Chicago, IL",
            "Feb 2024",
            "Present",
            (
                "Build FastAPI services for appointment scheduling that handle about 1.2 million requests per day.",
                "Designed idempotent payment-webhook processing with PostgreSQL row locks and a transactional outbox.",
                "Cut p95 latency of the provider search endpoint from 900 ms to 180 ms with query rewrites and composite indexes.",
                "Mentor two interns and maintain the team's pytest fixtures and CI pipeline.",
            ),
        ),
        Role(
            "Junior Backend Developer",
            "Brightline Commerce",
            "Chicago, IL",
            "Apr 2023",
            "Jan 2024",
            (
                "Developed Django REST Framework endpoints for the order management system.",
                "Moved slow report generation into Celery workers backed by Redis.",
                "Added integration tests with testcontainers and raised coverage from 48% to 81%.",
            ),
        ),
    ),
    skills=(
        ("Languages", "Python (asyncio, typing), SQL, Bash"),
        ("Frameworks", "FastAPI, Django, SQLAlchemy, Pydantic, Celery"),
        ("Data", "PostgreSQL, Redis"),
        ("Tooling", "Docker, GitHub Actions, pytest, testcontainers; Kubernetes (basic)"),
    ),
    education=("B.S. Software Engineering, Northfield Polytechnic",),
)

SOFIA = Resume(
    slug="sofia_martinez_frontend",
    name="Sofia Martinez",
    headline="Junior Frontend Engineer",
    location="Valencia, Spain",
    email="sofia.martinez@example.com",
    phone="+34 655 50 01 72",
    website="sofiamartinez.example.com",
    summary=(
        "Frontend developer with about a year and a half of professional experience building accessible React and "
        "TypeScript interfaces. Enjoys turning designs into polished, keyboard-friendly UI."
    ),
    roles=(
        Role(
            "Frontend Developer",
            "Lumen Learning Labs",
            "Remote",
            "Mar 2025",
            "Present",
            (
                "Build React and TypeScript components for a course-authoring tool used by 3,000 instructors.",
                "Raised the Lighthouse accessibility score from 72 to 96 by fixing labels, focus order and colour contrast.",
                "Wrote component tests with React Testing Library and introduced MSW for API mocking.",
            ),
        ),
        Role(
            "Frontend Intern",
            "Pixelwise Studio",
            "Valencia, Spain",
            "Jun 2024",
            "Dec 2024",
            ("Converted Figma designs into responsive marketing pages with Tailwind CSS.",),
        ),
    ),
    skills=(
        ("Languages", "JavaScript, TypeScript, HTML, CSS"),
        ("Frameworks", "React, Vite, Tailwind CSS"),
        ("Testing", "React Testing Library, Vitest, Playwright (basic)"),
        ("Other", "Git, Figma, WCAG 2.2"),
    ),
    projects=(
        ("Habit tracker PWA", "Offline-first React app using IndexedDB; about 1,200 monthly active users."),
    ),
    education=("B.A. Interactive Media, Coastal Arts University",),
)

PRIYA = Resume(
    slug="priya_raman_ml",
    name="Priya Raman",
    headline="Machine Learning Engineer (LLM)",
    location="Seattle, WA",
    email="priya.raman@example.com",
    phone="+1 (206) 555-0128",
    website=None,
    summary=(
        "Machine learning engineer with 5+ years across data science and ML engineering, now focused on "
        "retrieval-augmented generation, evaluation and low-latency LLM serving."
    ),
    roles=(
        Role(
            "Machine Learning Engineer",
            "Quillstack AI",
            "Remote",
            "May 2024",
            "Present",
            (
                "Built a RAG assistant over 2 million support articles with hybrid BM25 and embedding retrieval plus cross-encoder reranking; answer acceptance rose from 61% to 78%.",
                "Created an evaluation harness with 1,200 golden questions and an LLM judge calibrated against human labels.",
                "Served open-weight models with vLLM; reduced median time to first token from 1.4 s to 380 ms using prefix caching and continuous batching.",
            ),
        ),
        Role(
            "Data Scientist",
            "Meridian Insurance Analytics",
            "Hartford, CT",
            "Jul 2022",
            "Apr 2024",
            (
                "Built gradient-boosted fraud models (PR-AUC 0.41 to 0.58) deployed through an MLflow registry and a FastAPI service.",
                "Added drift monitoring with population stability index alerts.",
            ),
        ),
        Role(
            "Data Analyst",
            "Tidewater Energy",
            "Hartford, CT",
            "Jun 2021",
            "Jun 2022",
            ("Automated weekly forecasting reports with Python and SQL.",),
        ),
    ),
    skills=(
        ("ML & LLM", "PyTorch, scikit-learn, XGBoost, vLLM, sentence-transformers, FAISS, pgvector"),
        ("MLOps", "MLflow, Airflow, Docker, GitHub Actions"),
        ("Languages", "Python, SQL"),
    ),
    certifications=(),
    education=("M.S. Data Science, Riverbend University",),
)

ALL_RESUMES = (JANE, MARCUS, SOFIA, PRIYA)


# =============================================================================
# Plain text
# =============================================================================


def render_text(resume: Resume) -> str:
    lines = [
        resume.name.upper(),
        resume.headline,
        resume.contact_line,
        "",
        "SUMMARY",
        resume.summary,
        "",
        "EXPERIENCE",
    ]
    for role in resume.roles:
        lines += [f"{role.title} — {role.company}, {role.location}", f"{role.start} – {role.end}"]
        lines += [f"- {bullet}" for bullet in role.bullets]
        lines.append("")
    if resume.projects:
        lines.append("PROJECTS")
        lines += [f"- {name}: {desc}" for name, desc in resume.projects]
        lines.append("")
    lines.append("SKILLS")
    lines += [f"{group}: {items}" for group, items in resume.skills]
    lines.append("")
    if resume.certifications:
        lines.append("CERTIFICATIONS")
        lines += [f"- {c}" for c in resume.certifications]
        lines.append("")
    if resume.education:
        lines.append("EDUCATION")
        lines += [f"- {e}" for e in resume.education]
    return "\n".join(lines).rstrip() + "\n"


# =============================================================================
# PDF with a text layer (reportlab)
# =============================================================================


def render_pdf(resume: Resume, path: Path) -> None:
    accent = colors.HexColor("#1F4E79")
    styles = {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=accent),
        "headline": ParagraphStyle("headline", fontName="Helvetica", fontSize=12, leading=15),
        "contact": ParagraphStyle(
            "contact", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#444444")
        ),
        "section": ParagraphStyle(
            "section", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, textColor=accent
        ),
        "role": ParagraphStyle("role", fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=4),
        "dates": ParagraphStyle(
            "dates",
            fontName="Helvetica-Oblique",
            fontSize=9,
            leading=11,
            textColor=colors.HexColor("#555555"),
        ),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=9.5, leading=12.5, alignment=TA_LEFT),
        "bullet": ParagraphStyle(
            "bullet", fontName="Helvetica", fontSize=9.5, leading=12.5, leftIndent=12, bulletIndent=2
        ),
    }

    def section(title: str) -> list:  # type: ignore[type-arg]
        return [
            Paragraph(title.upper(), styles["section"]),
            HRFlowable(width="100%", thickness=0.6, color=accent, spaceBefore=1, spaceAfter=4),
        ]

    story: list = [  # type: ignore[type-arg]
        Paragraph(resume.name, styles["name"]),
        Paragraph(resume.headline, styles["headline"]),
        Paragraph(_xml(resume.contact_line), styles["contact"]),
        Spacer(1, 4),
        *section("Summary"),
        Paragraph(_xml(resume.summary), styles["body"]),
        *section("Experience"),
    ]
    for role in resume.roles:
        block = [
            Paragraph(_xml(f"{role.title} — {role.company}, {role.location}"), styles["role"]),
            Paragraph(f"{role.start} – {role.end}", styles["dates"]),
        ]
        block += [Paragraph(_xml(b), styles["bullet"], bulletText="•") for b in role.bullets]
        story.append(KeepTogether(block))
    if resume.projects:
        story += section("Projects")
        story += [
            Paragraph(f"<b>{_xml(n)}</b>: {_xml(d)}", styles["bullet"], bulletText="•")
            for n, d in resume.projects
        ]
    story += section("Skills")
    story += [Paragraph(f"<b>{_xml(g)}:</b> {_xml(items)}", styles["body"]) for g, items in resume.skills]
    if resume.certifications:
        story += section("Certifications")
        story += [Paragraph(_xml(c), styles["bullet"], bulletText="•") for c in resume.certifications]
    if resume.education:
        story += section("Education")
        story += [Paragraph(_xml(e), styles["body"]) for e in resume.education]

    doc = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title=f"{resume.name} — Resume",
        author=resume.name,
        subject=resume.headline,
        creator="sample resume generator",
        invariant=1,
    )
    doc.build(story)


def _xml(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# =============================================================================
# DOCX (python-docx)
# =============================================================================


def render_docx(resume: Resume, path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    accent = RGBColor(0x1F, 0x4E, 0x79)

    title = doc.add_paragraph()
    run = title.add_run(resume.name)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = accent
    headline = doc.add_paragraph()
    headline.add_run(resume.headline).font.size = Pt(12.5)
    contact = doc.add_paragraph()
    contact_run = contact.add_run(resume.contact_line)
    contact_run.font.size = Pt(9.5)
    contact.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def heading(text: str) -> None:
        h = doc.add_heading(text, level=2)
        for r in h.runs:
            r.font.color.rgb = accent

    heading("Summary")
    doc.add_paragraph(resume.summary)
    heading("Experience")
    for role in resume.roles:
        p = doc.add_paragraph()
        p.add_run(f"{role.title} — {role.company}").bold = True
        p.add_run(f", {role.location}")
        dates = doc.add_paragraph()
        dates_run = dates.add_run(f"{role.start} – {role.end}")
        dates_run.italic = True
        dates_run.font.size = Pt(9.5)
        for bullet in role.bullets:
            doc.add_paragraph(bullet, style="List Bullet")
    if resume.projects:
        heading("Projects")
        for name, desc in resume.projects:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{name}: ").bold = True
            p.add_run(desc)
    heading("Skills")
    for group, items in resume.skills:
        p = doc.add_paragraph()
        p.add_run(f"{group}: ").bold = True
        p.add_run(items)
    if resume.certifications:
        heading("Certifications")
        for cert in resume.certifications:
            doc.add_paragraph(cert, style="List Bullet")
    if resume.education:
        heading("Education")
        for edu in resume.education:
            doc.add_paragraph(edu)

    props = doc.core_properties
    props.author = resume.name
    props.title = f"{resume.name} — Resume"
    props.subject = resume.headline
    props.last_modified_by = resume.name
    props.created = DOC_DATE.replace(tzinfo=None)
    props.modified = DOC_DATE.replace(tzinfo=None)
    props.revision = 1
    doc.save(str(path))


# =============================================================================
# Images (Pillow): image resume and scanned PDF
# =============================================================================

_FONT_DIRS = (
    Path("/usr/share/fonts/truetype/dejavu"),
    Path("/usr/share/fonts/truetype/liberation"),
    Path("/Library/Fonts"),
)
_FONT_FILES = {
    "regular": ("DejaVuSans.ttf", "LiberationSans-Regular.ttf", "Arial.ttf"),
    "bold": ("DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf", "Arial Bold.ttf"),
}


def _font(kind: str, size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for directory in _FONT_DIRS:
        for name in _FONT_FILES[kind]:
            candidate = directory / name
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default(size=size)


def _wrap(text: str, font: ImageFont.FreeTypeFont | ImageFont.ImageFont, width: int) -> list[str]:
    words, lines, current = text.split(), [], ""
    for word in words:
        trial = f"{current} {word}".strip()
        if font.getlength(trial) <= width:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


@dataclass(frozen=True)
class _Line:
    text: str
    font: ImageFont.FreeTypeFont | ImageFont.ImageFont
    color: tuple[int, int, int]
    indent: int = 0
    gap_before: int = 0
    rule: bool = False


def _layout(resume: Resume) -> list[_Line]:
    accent, ink, grey = (31, 78, 121), (25, 25, 25), (85, 85, 85)
    f_name, f_head = _font("bold", 40), _font("regular", 25)
    f_section, f_role = _font("bold", 24), _font("bold", 20)
    f_body, f_small = _font("regular", 19), _font("regular", 17)
    width = PAGE_W - 2 * MARGIN
    out: list[_Line] = [_Line(resume.name, f_name, accent), _Line(resume.headline, f_head, ink, gap_before=4)]
    out += [_Line(t, f_small, grey, gap_before=6) for t in _wrap(resume.contact_line, f_small, width)]

    def section(title: str) -> None:
        out.append(_Line(title.upper(), f_section, accent, gap_before=24, rule=True))

    def paragraph(
        text: str, *, font=f_body, color=ink, indent: int = 0, bullet: bool = False, gap: int = 4
    ) -> None:  # type: ignore[no-untyped-def]
        wrapped = _wrap(text, font, width - indent - (22 if bullet else 0))
        for i, line in enumerate(wrapped):
            prefix = ("• " if i == 0 else "  ") if bullet else ""
            out.append(_Line(prefix + line, font, color, indent=indent, gap_before=gap if i == 0 else 0))

    section("Summary")
    paragraph(resume.summary)
    section("Experience")
    for role in resume.roles:
        paragraph(f"{role.title} — {role.company}, {role.location}", font=f_role, gap=14)
        paragraph(f"{role.start} – {role.end}", font=f_small, color=grey, gap=2)
        for bullet in role.bullets:
            paragraph(bullet, indent=10, bullet=True)
    if resume.projects:
        section("Projects")
        for name, desc in resume.projects:
            paragraph(f"{name}: {desc}", indent=10, bullet=True)
    section("Skills")
    for group, items in resume.skills:
        paragraph(f"{group}: {items}")
    if resume.certifications:
        section("Certifications")
        for cert in resume.certifications:
            paragraph(cert, indent=10, bullet=True)
    if resume.education:
        section("Education")
        for edu in resume.education:
            paragraph(edu)
    return out


def _line_height(line: _Line) -> int:
    ascent, descent = line.font.getmetrics() if hasattr(line.font, "getmetrics") else (16, 4)
    return int((ascent + descent) * 1.18) + line.gap_before + (8 if line.rule else 0)


def _draw_lines(lines: Sequence[_Line], height: int) -> Image.Image:
    image = Image.new("RGB", (PAGE_W, height), "white")
    draw = ImageDraw.Draw(image)
    y = MARGIN
    for line in lines:
        y += line.gap_before
        draw.text((MARGIN + line.indent, y), line.text, font=line.font, fill=line.color)
        ascent, descent = line.font.getmetrics() if hasattr(line.font, "getmetrics") else (16, 4)
        if line.rule:
            rule_y = y + ascent + descent + 4
            draw.line((MARGIN, rule_y, PAGE_W - MARGIN, rule_y), fill=line.color, width=2)
            y += 8
        y += int((ascent + descent) * 1.18)
    return image


def render_image(resume: Resume) -> Image.Image:
    """One tall page holding the whole resume (a typical "image resume")."""
    lines = _layout(resume)
    height = MARGIN * 2 + sum(_line_height(line) for line in lines)
    return _draw_lines(lines, max(height, PAGE_H))


def render_pages(resume: Resume) -> list[Image.Image]:
    """Letter-size pages at 150 DPI, breaking between lines."""
    pages: list[Image.Image] = []
    current: list[_Line] = []
    used = 0
    budget = PAGE_H - 2 * MARGIN
    for line in _layout(resume):
        needed = _line_height(line)
        if current and used + needed > budget:
            pages.append(_draw_lines(current, PAGE_H))
            current, used = [], 0
            line = _Line(line.text, line.font, line.color, line.indent, 0, line.rule)
            needed = _line_height(line)
        current.append(line)
        used += needed
    if current:
        pages.append(_draw_lines(current, PAGE_H))
    return pages


def _scan_effect(page: Image.Image, *, angle: float, seed: int) -> Image.Image:
    """Make a clean page look photocopied: greyscale, skew, blur, noise, speckles, edge shadow."""
    rng = np.random.default_rng(seed)
    grey = page.convert("L")
    rotated = grey.rotate(angle, resample=Image.Resampling.BICUBIC, expand=False, fillcolor=255)
    blurred = rotated.filter(ImageFilter.GaussianBlur(radius=0.7))
    pixels = np.asarray(blurred, dtype=np.float32)
    pixels = pixels * 0.93 + 8  # slightly lower contrast, off-white paper
    pixels += rng.normal(0.0, 9.0, size=pixels.shape)
    speckles = rng.random(pixels.shape) < 0.0009
    pixels[speckles] = rng.uniform(30, 120, size=int(speckles.sum()))
    shadow = np.linspace(55, 0, 40, dtype=np.float32)
    pixels[:, : shadow.size] -= shadow  # scanner lid shadow on the left edge
    return Image.fromarray(np.clip(pixels, 0, 255).astype(np.uint8), mode="L")


def render_scanned_pdf(resume: Resume, path: Path) -> int:
    angles = (-0.9, 0.7, -0.5, 0.4)
    pages = [
        _scan_effect(p, angle=angles[i % len(angles)], seed=1000 + i)
        for i, p in enumerate(render_pages(resume))
    ]
    pages[0].save(
        str(path),
        "PDF",
        save_all=True,
        append_images=pages[1:],
        resolution=float(DPI),
        title=f"{resume.name} — scanned resume",
        author="Scanner",
        creationDate=DOC_DATE.utctimetuple(),
        modDate=DOC_DATE.utctimetuple(),
    )
    return len(pages)


# =============================================================================
# Main
# =============================================================================


def generate(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    def note(path: Path) -> None:
        written.append(path)
        print(f"  {path.name:<32} {path.stat().st_size:>9,} bytes")

    # Jane Doe in every supported format.
    render_pdf(JANE, out_dir / f"{JANE.slug}.pdf")
    note(out_dir / f"{JANE.slug}.pdf")
    render_docx(JANE, out_dir / f"{JANE.slug}.docx")
    note(out_dir / f"{JANE.slug}.docx")
    (out_dir / f"{JANE.slug}.txt").write_text(render_text(JANE), encoding="utf-8")
    note(out_dir / f"{JANE.slug}.txt")
    image = render_image(JANE)
    image.save(out_dir / f"{JANE.slug}.png", optimize=True)
    note(out_dir / f"{JANE.slug}.png")
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=88, optimize=True)
    (out_dir / f"{JANE.slug}.jpg").write_bytes(buffer.getvalue())
    note(out_dir / f"{JANE.slug}.jpg")
    pages = render_scanned_pdf(JANE, out_dir / f"{JANE.slug}_scanned.pdf")
    note(out_dir / f"{JANE.slug}_scanned.pdf")
    print(f"    scanned PDF pages: {pages}")

    # Additional candidates, one format each.
    render_pdf(MARCUS, out_dir / f"{MARCUS.slug}.pdf")
    note(out_dir / f"{MARCUS.slug}.pdf")
    render_docx(SOFIA, out_dir / f"{SOFIA.slug}.docx")
    note(out_dir / f"{SOFIA.slug}.docx")
    (out_dir / f"{PRIYA.slug}.txt").write_text(render_text(PRIYA), encoding="utf-8")
    note(out_dir / f"{PRIYA.slug}.txt")
    return written


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--out", type=Path, default=DEFAULT_OUT, help=f"output directory (default: {DEFAULT_OUT})"
    )
    args = parser.parse_args(list(argv) if argv is not None else None)
    print(f"Writing sample resumes to {args.out}")
    generate(args.out.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
